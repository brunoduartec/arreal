from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "output/caderno_poderes_garoto/Transcricao_Caderno_de_Habilidades.docx"

AZUL = RGBColor(31, 78, 121)
AZUL_ESCURO = RGBColor(27, 52, 73)
CINZA = RGBColor(90, 90, 90)
DOURADO = RGBColor(157, 120, 47)


habilidades = [
    (
        1,
        "Volitação",
        [
            "A altura depende muito da confiança, pois não estamos [trecho incerto] medo.",
            "Em geral não é muito rápido.",
            "A sensação é de estar mais leve, como nadando.",
            "Inicia como sentindo que algo o puxa pra cima.",
        ],
    ),
    (
        2,
        "Invisibilidade",
        [
            "Você vê também através de você.",
            "A sensação é como se estivesse caindo, como se algo pressionasse sua cabeça o tempo todo.",
            "Durante a invisibilidade não se pode respirar, o que restringe o tempo de uso drasticamente, e a pessoa não pode ser tocada por outros.",
        ],
    ),
    (
        3,
        "Atravessar objetos",
        [
            "É preciso haver um contato inicial para reconhecer a substância, sendo o contato e o tato fundamentais; sentir como se fosse feito da mesma substância.",
            "Na verdade, o que acontece é que o corpo aos poucos se transforma na substância, mas só a área de contato.",
            "Perigoso se assustar e ficar preso ou perder uma parte do corpo.",
        ],
    ),
    (
        4,
        "Telecinese",
        [
            "Acontece por uma extensão astral e energética do próprio corpo na realidade, portanto só é possível fazer o que faria com a sua carne. Levantar objetos que suportaria.",
            "Inicia por ver o objeto e imaginar-se encostando nele; após uma conexão com o mesmo, se perder o contato visual acaba a conexão.",
        ],
    ),
    (
        5,
        "Guardar objetos no corpo",
        [
            "O princípio é parecido com atravessar objetos, mas também há uma redução do objeto.",
            "Acontece ao contrário: o objeto se torna corpo.",
            "Enquanto o objeto estiver em si, a pessoa se sente um pouco mais pesada.",
            "Não é uma adição integral de peso, mas a pessoa se sente mais pesada.",
            "Quanto mais treinada, menos peso sente, pois na real o objeto virou energia.",
        ],
    ),
    (
        6,
        "Viagem no tempo",
        [
            "Acontece com muita troca de energia, portanto não é sempre que pode ser usada.",
            "Uma vez utilizada, fica a pessoa cansada como se houvesse corrido uma maratona.",
            "Assim como o teletransporte, inicia por se imaginar em um local, salvo que a viagem no tempo precisa se enxergar também no tempo.",
            "Demanda um bom tempo de concentração para transportar-se para o espaço/tempo.",
            "Perigoso aparecer em um local que não mais existe ou dentro de algo.",
        ],
    ),
    (
        7,
        "Visão do tempo",
        [
            "Sente o que vai acontecer, mas como o presente não muda muito.",
            "Acontece uma inicial desconexão com o agora.",
            "Dependendo do nível de domínio da habilidade, a pessoa não consegue fazer outra coisa ao mesmo tempo.",
        ],
    ),
    (
        8,
        "Manipulação de energia",
        [
            "Formada uma esfera de energia entre as mãos.",
            "É preciso sentir a energia fluindo pelo corpo até chegar nas mãos.",
            "Enquanto estiver sentindo a conexão com a energia, ela existe.",
            "Pode causar danos, pois ao entrar em contato com outra matéria, esta é desestabilizada.",
        ],
    ),
    (
        9,
        "Teletransporte",
        [
            "Inicia por começar a sentir-se no outro local, formando mentalmente a ideia do outro local.",
            "Em geral a sensação inicia na nuca e, enquanto não completar o transporte, é imprescindível não perder a concentração no local de destino.",
        ],
    ),
    (
        11,
        "Influência mental",
        [
            "Tal habilidade tem 2 vertentes: leitura mental e influência.",
            "Na verdade, o que acontece é a conexão mental entre as pessoas, e as influências se dão ao fato de as pessoas serem influenciáveis.",
            "O que acontece é um sugestionamento que se inicia por tentar sentir como o outro está e sincronizar-se um com o outro.",
            "Tal habilidade só se dá por uma conexão visual ou mesmo auditiva.",
        ],
    ),
    (
        13,
        "Cura",
        [
            "As curas ocorrem por manipulação de energia pelo pensamento.",
            "O processo inicia por visualizar a parte afetada e imaginar energia envolvendo a parte; parte por sentir como se tocasse a parte.",
            "O mesmo tem que ser feito com muita atenção e foco; se não for concluído completamente, se perde.",
        ],
    ),
]


def set_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    set_font(run, 9, color=CINZA)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

h1 = styles["Heading 1"]
h1.font.name = "Calibri"
h1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
h1.font.size = Pt(16)
h1.font.bold = True
h1.font.color.rgb = AZUL
h1.paragraph_format.space_before = Pt(18)
h1.paragraph_format.space_after = Pt(10)
h1.paragraph_format.keep_with_next = True

h2 = styles["Heading 2"]
h2.font.name = "Calibri"
h2._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
h2.font.size = Pt(13)
h2.font.bold = True
h2.font.color.rgb = AZUL
h2.paragraph_format.space_before = Pt(14)
h2.paragraph_format.space_after = Pt(7)
h2.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(header.add_run("ARREAL  |  CADERNO DE HABILIDADES"), 9, True, color=CINZA)
add_page_number(section.footer.paragraphs[0])

for _ in range(5):
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

kicker = doc.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(kicker.add_run("TRANSCRIÇÃO DAS ANOTAÇÕES ORIGINAIS"), 10, True, color=DOURADO)
kicker.paragraph_format.space_after = Pt(18)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(title.add_run("Caderno de Habilidades"), 28, True, color=AZUL_ESCURO)
title.paragraph_format.space_after = Pt(8)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(subtitle.add_run("Texto extraído das fotografias manuscritas"), 14, italic=True, color=AZUL)
subtitle.paragraph_format.space_after = Pt(34)

note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(
    note.add_run(
        "Transcrição conservadora. Pontuação e acentuação foram ajustadas apenas para leitura. "
        "Trechos duvidosos estão marcados; nenhuma informação foi completada por fontes secundárias."
    ),
    10,
    color=CINZA,
)

doc.add_page_break()

doc.add_heading("Critérios da transcrição", level=1)
p = doc.add_paragraph()
set_font(p.add_run("Fonte principal: "), bold=True, color=AZUL_ESCURO)
set_font(p.add_run("as onze fotografias enviadas nesta conversa."))

p = doc.add_paragraph()
set_font(p.add_run("Ordem: "), bold=True, color=AZUL_ESCURO)
set_font(p.add_run("numeração manuscrita das habilidades, não a ordem de envio das imagens."))

p = doc.add_paragraph()
set_font(p.add_run("Lacunas: "), bold=True, color=AZUL_ESCURO)
set_font(p.add_run("não foram recebidas páginas correspondentes às habilidades nº 10 e nº 12."))

doc.add_heading("Transcrição", level=1)

for numero, nome, paragrafos in habilidades:
    doc.add_heading(f"{numero}. {nome}", level=2)
    for texto in paragrafos:
        p = doc.add_paragraph(texto)
        p.paragraph_format.keep_together = True

doc.add_heading("Páginas não recebidas", level=1)
doc.add_paragraph(
    "As habilidades nº 10 e nº 12 aparecem como lacunas na sequência numérica. "
    "Se essas fotografias forem enviadas posteriormente, devem ser inseridas aqui sem alterar as demais transcrições."
)

doc.core_properties.title = "Caderno de Habilidades — Transcrição"
doc.core_properties.subject = "Transcrição das anotações manuscritas enviadas em fotografias"
doc.core_properties.author = "Bruno Duarte Corrêa"
doc.save(OUTPUT)
print(OUTPUT)
